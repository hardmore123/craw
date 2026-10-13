"""输入：报备清单明细表；输出两份「沿用原表布局」的文件：
  1) 抓取结果导出.xlsx —— 制造商/型号/标准 列换成官网抓取值（蓝字），M=查询状态，N=来源URL
  2) 对比结果导出.xlsx —— 表格原值不动，M=核对结论（一致/不一致/未查到 + 哪些字段），不一致红字，N=来源URL

特性：
  - 自动定位输入 xlsx（文件名可含日期后缀，无需改脚本）；
  - 按表头自动识别列（证书号/认证类别/型号/制造商/标准），CB、UL、CB-UL 三表通吃；
  - 一次抓取、同时产出两份，避免 UL 重复抓。

用法：python 导出两表.py [每表条数] [sheet1,sheet2,...]   (0=全部)
"""
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from cert_verify.config import ColumnMap, RunConfig  # noqa: E402
from cert_verify.verifier import VerifyEngine, summarize_outcomes  # noqa: E402
from cert_verify.excel_writer import write_crawl_copy, write_results  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))


def find_input() -> str:
    cands = [f for f in os.listdir(HERE)
             if f.endswith(".xlsx") and "报备" in f and not f.startswith("~$")
             and "结果" not in f]
    if not cands:
        raise SystemExit("在 本地Windows/交付/input 下没找到报备清单 xlsx（文件名需含“报备”）")
    cands.sort(key=lambda f: os.path.getmtime(os.path.join(HERE, f)), reverse=True)
    return os.path.join(HERE, cands[0])


def detect_columns(ws) -> ColumnMap:
    """读第 1 行表头，按关键词自动定位各列字母。"""
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v is None:
            continue
        hdr[get_column_letter(c)] = str(v).replace("\n", " ")

    def find(*keys, default=""):
        for col, text in hdr.items():
            if any(k in text for k in keys):
                return col
        return default

    cert = find("证书号", "Mark(s)")
    agency = find("认证类别")
    model = find("型号", "Model")
    mfr = find("制造商", "商标", "Manufacturer")
    std = find("认证标准", "标准", "Standard")
    # 结果/详情列：优先用空列，默认为 M/N（报备清单一般 M、N 为空）
    return ColumnMap(cert_no=cert or "I", result="M", manufacturer=mfr, model=model,
                     standard=std, cert_status="", factory="", agency=agency,
                     first_data_row=2)


XLSX = find_input()
OUT_CRAWL = os.path.join(HERE, "抓取结果导出.xlsx")
OUT_DIFF = os.path.join(HERE, "对比结果导出.xlsx")
PROXY = os.environ.get("CERT_PROXY", "http://127.0.0.1:7877")

limit = int(sys.argv[1]) if len(sys.argv) > 1 else 0
sheets = sys.argv[2].split(",") if len(sys.argv) > 2 else ["认证类别CB", "认证类别UL", "认证类别CB-UL"]

print(f"输入文件：{os.path.basename(XLSX)}", flush=True)
wb_probe = load_workbook(XLSX, read_only=True, data_only=True)
colmap_by_sheet = {}
for sh in sheets:
    if sh in wb_probe.sheetnames:
        cm = detect_columns(wb_probe[sh])
        colmap_by_sheet[sh] = cm
        print(f"  [{sh}] 列映射: 证书号={cm.cert_no} 认证类别={cm.agency} "
              f"型号={cm.model} 制造商={cm.manufacturer} 标准={cm.standard}", flush=True)
    else:
        print(f"  [warn] 无此 sheet: {sh}", flush=True)
wb_probe.close()

outcomes_by_sheet = {}
t0 = time.time()
for sh, cm in colmap_by_sheet.items():
    cfg = RunConfig(xlsx_path=XLSX, sheets=[sh], columns=cm, limit=limit,
                    row_delay=0.1, stop_on_blocked=False, proxy=PROXY)
    print(f"\n########## 抓取 sheet={sh} limit={limit} ##########", flush=True)
    lst = []
    for o in VerifyEngine(cfg).iter_rows():
        lst.append(o)
        print(f"  r{o.row:<4} {o.cert_no:<18} -> {o.agency:<13} [{o.status}]  (+{time.time()-t0:.0f}s)",
              flush=True)
    outcomes_by_sheet[sh] = lst
    print("  " + str(summarize_outcomes(lst)), flush=True)

print("\n写出两份文件...", flush=True)
i1 = write_crawl_copy(XLSX, OUT_CRAWL, outcomes_by_sheet, detail_col="N")
i2 = write_results(XLSX, OUT_DIFF, outcomes_by_sheet, detail_col="N")
print(f"  抓取结果导出 -> {OUT_CRAWL}  {i1.get('message') or 'OK'}")
print(f"  对比结果导出 -> {OUT_DIFF}  {i2.get('message') or 'OK'}")

allo = [o for lst in outcomes_by_sheet.values() for o in lst]
print("\n总计 " + str(summarize_outcomes(allo)))
