"""通用网评汇总模板：多站点/多地区共用的一行一评价宽表。

列顺序在现有 Amazon 网评表（15 列）基础上，扩展 AI 处理列与溯源/价格列，
使 Amazon（加拿大等）与 価格.com（日本）等站点共用同一套模板：

  1  所属国家
  2  品牌
  3  型号
  4  渠道
  5  尺寸
  6  网评分数(星级)
  7  网评标题
  8  网评内容
  9  翻译              （借助 AI/翻译引擎，未接入时留空）
  10 提炼优点(通过AI)   （未接入时留空）
  11 提炼缺点(通过AI)   （未接入时留空）
  12 赞同数
  13 评论时间
  14 评论链接
  15 图片数量
  16 图片地址
  17 是否有视频
  18 视频地址
  19 商品ID            （站内唯一商品号，如 kakaku item / Amazon ASIN）
  20 价格              （用于后续价格爬取/比价，可留空）

翻译/优缺点三列紧跟「网评内容」之后，便于逐条对照原文。这三列的值由可插拔
钩子产生（见 AiHooks），默认返回空字符串。历史 Amazon 15 列文件的字段仍可
按字段名对齐（列位置变化不影响按表头读取）。
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

COLUMNS = [
    "所属国家", "品牌", "型号", "渠道", "尺寸",
    "网评分数(星级)", "网评标题", "网评内容",
    "翻译", "提炼优点(通过AI)", "提炼缺点(通过AI)",
    "赞同数", "评论时间", "评论链接", "图片数量", "图片地址",
    "是否有视频", "视频地址", "商品ID", "价格",
]


@dataclass
class ReviewRecord:
    """一条网评的通用记录，字段与模板列一一对应。"""
    country: str = ""
    brand: str = ""
    model: str = ""
    channel: str = ""
    size: str = ""
    rating: Any = ""
    title: str = ""
    body: str = ""
    helpful: Any = ""
    review_time: str = ""
    review_url: str = ""
    image_count: Any = 0
    image_urls: str = ""
    has_video: str = "否"
    video_urls: str = ""
    translation: str = ""
    ai_pros: str = ""
    ai_cons: str = ""
    item_id: str = ""
    price: str = ""

    def as_row(self) -> list[Any]:
        return [
            self.country, self.brand, self.model, self.channel, self.size,
            self.rating, self.title, self.body,
            self.translation, self.ai_pros, self.ai_cons,
            self.helpful, self.review_time, self.review_url, self.image_count,
            self.image_urls, self.has_video, self.video_urls,
            self.item_id, self.price,
        ]


# ==================== AI / 翻译 可插拔钩子（默认不实现）====================

@dataclass
class AiHooks:
    """AI 处理钩子集合，默认全部返回空字符串（不接入任何服务）。

    将来接入方式（任选其一，只改这里，不动抓取与导出）：
      - 在线免费翻译引擎：deep-translator 的 GoogleTranslator 等（需联网，可能限流）；
      - 本地离线翻译：NLLB/argos-translate（不联网，需安装对应模型）；
      - 大模型：调用自有 OpenAI-compatible 服务做翻译 + 优缺点提炼。
    传入 AiHooks(translate=fn, extract_pros=fn, extract_cons=fn) 即可启用。
    """
    translate: Callable[[str], str] | None = None
    extract_pros: Callable[[str], str] | None = None
    extract_cons: Callable[[str], str] | None = None

    def apply(self, rec: ReviewRecord) -> None:
        if self.translate and not rec.translation:
            try:
                rec.translation = self.translate(rec.body) or ""
            except Exception:
                rec.translation = ""
        if self.extract_pros and not rec.ai_pros:
            try:
                rec.ai_pros = self.extract_pros(rec.body) or ""
            except Exception:
                rec.ai_pros = ""
        if self.extract_cons and not rec.ai_cons:
            try:
                rec.ai_cons = self.extract_cons(rec.body) or ""
            except Exception:
                rec.ai_cons = ""


NULL_HOOKS = AiHooks()          # 默认：三列全空


# ==================== 导出 ====================

def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def write_csv(path: str | Path, records: list[ReviewRecord],
              hooks: AiHooks = NULL_HOOKS) -> int:
    path = Path(path)
    _ensure_parent(path)
    n = 0
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        for rec in records:
            hooks.apply(rec)
            writer.writerow(rec.as_row())
            n += 1
    return n


def write_xlsx(path: str | Path, records: list[ReviewRecord],
               hooks: AiHooks = NULL_HOOKS, sheet_title: str = "网评汇总") -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    _ensure_parent(path)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(COLUMNS)
    fill = PatternFill("solid", fgColor="E38B29")      # 与旧 Amazon 表一致的橙色表头
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    n = 0
    for rec in records:
        hooks.apply(rec)
        row = rec.as_row()
        ws.append(row)
        # openpyxl 会把以“=”开头的翻译识别成公式；所有字符串列强制按文本保存，
        # 避免用户名或正文导致 XLSX 读取为空或被 Excel 执行。
        for cell, value in zip(ws[ws.max_row], row):
            if isinstance(value, str):
                cell.data_type = "s"
        n += 1
    widths = [10, 10, 16, 12, 8, 12, 22, 50,
              40, 24, 24,
              8, 20, 42, 8, 30, 8, 24, 14, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    return n
