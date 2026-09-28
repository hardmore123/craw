"""加拿大线价格监控汇总模板：一行一机型，每个零售站一列价格。

与日本线（price_export.py，kakaku 单站 8 个店铺列）不同，加拿大线的价格来自多个
零售站（Amazon / Best Buy / Walmart / Costco / The Brick / Leon's / Visions /
Canadian Tire），因此列布局为「固定头 + 每个零售站一列价格 + 溯源尾」。零售站列
从 regions_config 的 ca.retail_sites 取（保持与抓取端一致，新增/调整零售站只改配置）。

列顺序：
  1  所属国家
  2  品牌
  3  机型
  4  尺寸
  5  货币                 （CAD）
  6  最低价               （各零售站中的最低售价）
  7  最低价零售站         （最低价对应的零售站，溯源）
  8..  各零售站价格列    （regions_config ca.retail_sites 顺序）
  --- 溯源尾 ---
  采集时间

价格与网评在同一商品页同步抓取（见 sites 零售适配器 parse_product），价格线只取
其中的 PriceSnapshot 部分汇总到本表。
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .regions_config import get_region


def retail_columns(region_code: str = "ca") -> list[tuple[str, str]]:
    """返回该区域零售站的 (site_code, 展示名) 列表，作为价格列。"""
    region = get_region(region_code)
    if not region:
        return []
    return [(r.site, r.name) for r in region.retail_sites]


FIXED_HEAD = ["所属国家", "品牌", "机型", "尺寸", "货币", "最低价", "最低价零售站"]
FIXED_TAIL = ["采集时间"]


def build_columns(region_code: str = "ca") -> tuple[list[str], list[str]]:
    """构造 (内部列 key, 对外展示表头)。零售站列 key=site_code，展示=零售站名。"""
    cols = retail_columns(region_code)
    internal = FIXED_HEAD + [site for site, _ in cols] + FIXED_TAIL
    display = FIXED_HEAD + [name for _, name in cols] + FIXED_TAIL
    return internal, display


@dataclass
class CaPriceRecord:
    """一个机型在某次采集时各零售站的价格快照。"""
    country: str = "Canada"
    brand: str = ""
    model: str = ""
    size: str = ""
    currency: str = "CAD"
    retail_prices: dict[str, Any] = field(default_factory=dict)   # {site_code: 价}
    captured_at: str = ""

    def lowest(self) -> tuple[Any, str]:
        """返回 (最低价, 最低价零售站 site_code)；无有效价返回 ('', '')。"""
        best_price = None
        best_site = ""
        for site, val in self.retail_prices.items():
            try:
                p = float(str(val).replace(",", "").replace("$", "").strip())
            except (TypeError, ValueError):
                continue
            if p <= 0:
                continue
            if best_price is None or p < best_price:
                best_price = p
                best_site = site
        return ("" if best_price is None else best_price), best_site

    def as_row(self, region_code: str = "ca") -> list[Any]:
        low, low_site = self.lowest()
        cols = retail_columns(region_code)
        low_name = next((name for site, name in cols if site == low_site), low_site)
        head = [self.country, self.brand, self.model, self.size,
                self.currency, low, low_name]
        prices = [self.retail_prices.get(site, "") for site, _ in cols]
        tail = [self.captured_at]
        return head + prices + tail


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def write_csv(path: str | Path, records: list[CaPriceRecord],
              region_code: str = "ca") -> int:
    path = Path(path)
    _ensure_parent(path)
    _, display = build_columns(region_code)
    n = 0
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(display)
        for rec in records:
            writer.writerow(rec.as_row(region_code))
            n += 1
    return n


def write_xlsx(path: str | Path, records: list[CaPriceRecord],
               region_code: str = "ca", sheet_title: str = "价格监控") -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    _ensure_parent(path)
    _, display = build_columns(region_code)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(display)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    n = 0
    for rec in records:
        row = rec.as_row(region_code)
        ws.append(row)
        for cell, value in zip(ws[ws.max_row], row):
            if isinstance(value, str):
                cell.data_type = "s"
        n += 1
    ws.freeze_panes = "A2"
    wb.save(path)
    return n
