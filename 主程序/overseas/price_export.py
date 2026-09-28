"""价格监控汇总表：一行一机型（或型号×零售站）的分列价格宽表。

默认（kakaku 日本线）列布局：
  1  所属区域              （按品牌来源站点划分，如 Japan）
  2  品牌
  3  机型
  4  尺寸
  5  上市时间
  6  货币                  （价格列的货币单位，kakaku 为 JPY）
  7  最低价                （= kakaku 最安価格，第一列标的价）
  8..15  8 个店铺列        （エディオンネストショップ / ヨドバシ / ... / Amazon）
  --- 溯源列 ---
  16 最安店铺              （最安価格对应的店铺名，便于核对第一列）
  17 商品ID                （kakaku item id，如 K0001785507）
  18 商品链接
  19 采集时间              （YYYY-MM-DD HH:MM）

**店铺列动态化（2026-09-17 新增，任务表 P0-3）**：
  传入 RetailSpec 时，店铺段按 spec 的 shops 段动态生成：
    - shops.mode=multi_shop → 读 spec.shops.columns（聚合站，如 kakaku 的 8 家）
    - shops.mode=single_shop → 简化两列「渠道 / 现价」（自营站如 Amazon/Liverpool）
    - 不传 spec（向后兼容） → 回落日本 8 个店铺列常量
  调用方传 spec 走 build_columns(spec)；产出的行用 PriceRecord.as_row(spec)。
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

# 模板 8 个店铺列（顺序即导出列顺序），与适配器 SHOP_COLUMNS 的列名一致。
# 这些是数据字典的 key（内部对齐用），不要改动；对外展示的中文表头见 SHOP_COLUMN_ZH。
SHOP_COLUMN_NAMES = [
    "エディオンネットショップ",
    "ヨドバシ.com",
    "ヤマダウェブコム",
    "ビックカメラ.com",
    "ケーズデンキ",
    "Joshin",
    "ノジマオンライン",
    "Amazon",
]

# 店铺列的中文显示名（仅用于导出表头，不改变内部数据 key）。
# 采用店铺品牌的通行中文译法，保留其英文商号以便核对。
SHOP_COLUMN_ZH = {
    "エディオンネットショップ": "爱电王网络商店(EDION)",
    "ヨドバシ.com": "友都八喜(Yodobashi.com)",
    "ヤマダウェブコム": "山田电机网购(Yamada)",
    "ビックカメラ.com": "BIC CAMERA.com",
    "ケーズデンキ": "K's电器(K's Denki)",
    "Joshin": "上新电机(Joshin)",
    "ノジマオンライン": "Nojima 在线",
    "Amazon": "亚马逊(Amazon)",
}

FIXED_HEAD = ["所属区域", "品牌", "机型", "尺寸", "上市时间", "货币", "最低价"]
FIXED_TAIL = ["最安店铺", "商品ID", "商品链接", "采集时间"]

# single_shop（自营站）简化列：无店铺列，改「渠道 / 现价」两列（任务表 P0-3 3.3）
_SINGLE_SHOP_KEYS = ["渠道", "现价"]


def shop_column_names(spec=None) -> list[str]:
    """返回价格表的「中间段」列（内部 key）。

    - spec 为 multi_shop  → 读 spec.shops.columns 的 key（该站声明的店铺列）
    - spec 为 single_shop → ['渠道', '现价']（简化列）
    - 不传 spec           → 日本 kakaku 8 店铺列（向后兼容）
    """
    if spec is None:
        return list(SHOP_COLUMN_NAMES)
    shops = spec.get("shops") or {}
    mode = shops.get("mode")
    if mode == "multi_shop":
        return [c["key"] for c in (shops.get("columns") or []) if c.get("key")]
    if mode == "single_shop":
        return list(_SINGLE_SHOP_KEYS)
    # 未知/缺失 mode：回落日本 8 列（与现状一致）
    return list(SHOP_COLUMN_NAMES)


def shop_column_display(spec=None) -> tuple:
    """价格表「店铺段」的对外中文表头，与 shop_column_names 顺序一致。"""
    names = shop_column_names(spec)
    if spec is None:
        return [SHOP_COLUMN_ZH.get(name, name) for name in names]
    shops = spec.get("shops") or {}
    mode = shops.get("mode")
    if mode == "multi_shop":
        display = {c.get("key"): c.get("display") or c.get("key")
                   for c in (shops.get("columns") or [])}
        return [display.get(name, name) for name in names]
    if mode == "single_shop":
        return list(_SINGLE_SHOP_KEYS)     # 简化列直接用中文表头
    return [SHOP_COLUMN_ZH.get(name, name) for name in names]


def build_columns(spec: dict = None) -> tuple:
    """构造 (内部列 key, 对外展示表头)。店铺段按 spec 动态生成（任务表 P0-3 3.1）。"""
    internal = FIXED_HEAD + shop_column_names(spec) + FIXED_TAIL
    display = FIXED_HEAD + shop_column_display(spec) + FIXED_TAIL
    return internal, display


# 与旧版命名兼容：原有调用方直接读这个常量时回落日本 8 列。
# 传 spec 的调用走 build_columns()。
COLUMNS, DISPLAY_COLUMNS = build_columns()


@dataclass
class PriceRecord:
    """一个机型在某次采集时的各店铺价格快照，字段与模板列一一对应。"""
    region: str = ""                       # 所属区域（按品牌来源站点，如 Japan）
    brand: str = ""
    model: str = ""
    size: str = ""
    release: str = ""
    currency: str = ""                     # 货币单位，如 JPY
    lowest_price: Any = ""                 # 最低价（multi_shop 为各店铺最低；single 为现价）
    shop_prices: dict[str, Any] = field(default_factory=dict)   # {店铺列名: 价}
    lowest_shop: str = ""                  # 最低价对应店铺名（溯源）
    channel: str = ""                      # single_shop 的渠道名（渠道即店铺）；多店铺留空
    item_id: str = ""
    url: str = ""
    captured_at: str = ""

    def as_row(self, spec: dict = None) -> list[Any]:
        """展开为一行。店铺段按 spec 动态取列（无 spec 回落日本 8 列）。"""
        head = [self.region, self.brand, self.model, self.size, self.release,
                self.currency, self.lowest_price]
        mode = None
        if spec is not None:
            mode = (spec.get("shops") or {}).get("mode")
        if mode == "multi_shop":
            shops_cells = [self.shop_prices.get(n, "") for n in shop_column_names(spec)]
        elif mode == "single_shop":
            shops_cells = [self.channel, self.lowest_price]
        else:
            shops_cells = [self.shop_prices.get(n, "") for n in SHOP_COLUMN_NAMES]
        tail = [self.lowest_shop, self.item_id, self.url, self.captured_at]
        return head + shops_cells + tail


# ==================== 分块排序 ====================

def _size_key(size: str) -> int:
    """尺寸排序键：从「65"」「100V」等取数字，取不到排最后。"""
    import re
    m = re.search(r"(\d{2,3})", size or "")
    return int(m.group(1)) if m else -1


def group_by_brand(records: list[PriceRecord]) -> list[tuple[str, str, list[PriceRecord]]]:
    """把记录按 (区域, 品牌) 聚成块，返回 [(区域, 品牌, 该块记录), ...]。

    块的先后按「首次出现顺序」保持稳定（与传入品牌顺序一致）；
    块内机型按尺寸从大到小、其次机型名排序，便于横向对齐同品牌不同尺寸。
    """
    order: list[tuple[str, str]] = []
    buckets: dict[tuple[str, str], list[PriceRecord]] = {}
    for rec in records:
        key = (rec.region or "", rec.brand or "")
        if key not in buckets:
            buckets[key] = []
            order.append(key)
        buckets[key].append(rec)
    blocks: list[tuple[str, str, list[PriceRecord]]] = []
    for region, brand in order:
        rows = sorted(buckets[(region, brand)],
                      key=lambda r: (-_size_key(r.size), r.model))
        blocks.append((region, brand, rows))
    return blocks


# ==================== 导出 ====================

def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def write_csv(path: str | Path, records: list[PriceRecord],
              blocked: bool = True, spec: dict = None) -> int:
    """导出 CSV。blocked=True 时按品牌分块，每块前插一行品牌分隔标题。

    spec 可选：传 RetailSpec 时店铺列动态化（multi_shop→columns；single_shop→渠道/现价）。
    """
    path = Path(path)
    _ensure_parent(path)
    _, display = build_columns(spec)
    ncol = len(display)
    n = 0
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(display)
        if not blocked:
            for rec in records:
                writer.writerow(rec.as_row(spec))
                n += 1
            return n
        for region, brand, rows in group_by_brand(records):
            # 品牌分隔标题行：第一列写「【区域 / 品牌】机型数」，其余留空
            title = f"【{region} / {brand}】{len(rows)} 机型" if region else f"【{brand}】{len(rows)} 机型"
            sep = [title] + [""] * (ncol - 1)
            writer.writerow(sep)
            for rec in rows:
                writer.writerow(rec.as_row(spec))
                n += 1
    return n


def write_xlsx(path: str | Path, records: list[PriceRecord],
               sheet_title: str = "价格监控", blocked: bool = True,
               spec: dict = None) -> int:
    """导出 XLSX。blocked=True 时按品牌分块，每块前加一行带底色的品牌标题。

    spec 可选：传 RetailSpec 时店铺列动态化（multi_shop→columns；single_shop→渠道/现价）。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    _ensure_parent(path)
    _, display = build_columns(spec)
    ncol = len(display)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    head_fill = PatternFill("solid", fgColor="E38B29")     # 表头橙色（与网评表一致）
    head_font = Font(color="FFFFFF", bold=True)
    block_fill = PatternFill("solid", fgColor="FCE4CC")    # 品牌分隔行浅橙
    block_font = Font(color="7A3E00", bold=True)

    ws.append(display)
    for cell in ws[1]:
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    n = 0

    def _append_data(rec: PriceRecord) -> None:
        ws.append(rec.as_row(spec))
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top")

    if blocked:
        for region, brand, rows in group_by_brand(records):
            title = (f"【{region} / {brand}】{len(rows)} 机型"
                     if region else f"【{brand}】{len(rows)} 机型")
            ws.append([title] + [""] * (ncol - 1))
            trow = ws.max_row
            ws.merge_cells(start_row=trow, start_column=1, end_row=trow, end_column=ncol)
            tcell = ws.cell(row=trow, column=1)
            tcell.fill = block_fill
            tcell.font = block_font
            tcell.alignment = Alignment(horizontal="left", vertical="center")
            for rec in rows:
                _append_data(rec)
                n += 1
    else:
        for rec in records:
            _append_data(rec)
            n += 1

    # 列宽：区域/品牌/机型/尺寸/上市/货币/最低价 + 动态店铺段 + 溯源
    shop_w = len(display) - len(FIXED_HEAD) - len(FIXED_TAIL)
    widths = ([10, 10, 14, 8, 16, 8, 12]
              + [16] * shop_w
              + [16, 14, 44, 18])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "D2"                              # 冻结区域/品牌/机型 + 表头
    wb.save(path)
    return n
