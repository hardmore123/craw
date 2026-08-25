"""本地 HTTP 服务：静态页 + JSON API。仅绑 127.0.0.1。

零额外依赖（标准库 http.server）。后台线程跑核对，前端轮询进度。
"""
from __future__ import annotations

import json
import os
import re
import threading
import urllib.parse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from ..config import ColumnMap
from ..excel_runner import inspect_workbook
from .tasks import WORK_DIR, TaskManager

STATIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")
MAX_UPLOAD = 30 * 1024 * 1024        # 30MB
_UL_PROFILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                           ".ul_profile")

_mgr = TaskManager()
_ul_login_running = threading.Event()

_CONTENT_TYPES = {
    ".html": "text/html; charset=utf-8",
    ".js": "application/javascript; charset=utf-8",
    ".css": "text/css; charset=utf-8",
}


class Handler(BaseHTTPRequestHandler):
    server_version = "CertVerifyWeb/1.0"

    def log_message(self, *args):        # 静默默认访问日志
        pass

    # ---------------- 响应helpers ----------------
    def _json(self, obj, code=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _err(self, msg, code=400):
        self._json({"error": msg}, code)

    def _static(self, path):
        rel = "index.html" if path in ("/", "") else path.lstrip("/")
        full = os.path.normpath(os.path.join(STATIC_DIR, rel))
        if not full.startswith(STATIC_DIR) or not os.path.isfile(full):
            self._err("Not found", 404)
            return
        ext = os.path.splitext(full)[1].lower()
        with open(full, "rb") as f:
            body = f.read()
        self.send_response(200)
        self.send_header("Content-Type", _CONTENT_TYPES.get(ext, "application/octet-stream"))
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    # ---------------- GET ----------------
    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path, qs = parsed.path, urllib.parse.parse_qs(parsed.query)
        try:
            if path == "/api/progress":
                return self._progress(qs)
            if path == "/api/ul-status":
                return self._ul_status()
            if path == "/api/download":
                return self._download(qs)
            return self._static(path)
        except Exception as e:  # noqa: BLE001
            self._err(f"{type(e).__name__}: {e}", 500)

    # ---------------- POST ----------------
    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        try:
            if path == "/api/upload":
                return self._upload()
            if path == "/api/run":
                return self._run()
            if path == "/api/cancel":
                return self._cancel()
            if path == "/api/ul-login":
                return self._ul_login()
            self._err("Not found", 404)
        except Exception as e:  # noqa: BLE001
            self._err(f"{type(e).__name__}: {e}", 500)

    # ---------------- API: upload ----------------
    def _upload(self):
        ctype = self.headers.get("Content-Type", "")
        clen = int(self.headers.get("Content-Length", 0))
        if clen <= 0 or clen > MAX_UPLOAD:
            return self._err(f"文件为空或超过 {MAX_UPLOAD // 1024 // 1024}MB 上限")
        if "multipart/form-data" not in ctype:
            return self._err("需以 multipart/form-data 上传")
        m = re.search(r"boundary=([^;]+)", ctype)
        if not m:
            return self._err("缺少 multipart boundary")
        boundary = m.group(1).strip('"').encode()
        raw = self.rfile.read(clen)
        filename, filedata = _parse_multipart_file(raw, boundary)
        if not filedata:
            return self._err("未找到上传文件")
        if not filename.lower().endswith(".xlsx"):
            return self._err("只接受 .xlsx 文件")

        file_id = os.urandom(6).hex()
        d = os.path.join(WORK_DIR, file_id)
        os.makedirs(d, exist_ok=True)
        safe_name = os.path.basename(filename)

        # 加密检测 + 解密兜底
        from ..decrypt import is_encrypted, decrypt_file, DECRYPT_ENABLED
        if is_encrypted(filedata):
            if not DECRYPT_ENABLED:
                return self._err(
                    "文件可能被加密（不是标准 xlsx 格式）。"
                    "请在本机用 Excel 打开后「另存为 .xlsx」再上传，"
                    "或联系管理员启用解密服务。")
            try:
                filedata = decrypt_file(filedata, safe_name)
            except RuntimeError as e:
                return self._err(
                    f"文件被加密，自动解密失败：{e}。"
                    "请在本机用 Excel 打开后「另存为 .xlsx」再上传。")

        dst = os.path.join(d, safe_name)
        with open(dst, "wb") as f:
            f.write(filedata)

        try:
            sheets = _read_sheets_preview(dst)
        except Exception as e:  # noqa: BLE001
            return self._err(f"无法解析 xlsx：{type(e).__name__}: {e}")
        return self._json({"file_id": file_id, "filename": safe_name, "sheets": sheets})

    # ---------------- API: run ----------------
    def _run(self):
        data = self._read_json()
        file_id = data.get("file_id", "")
        src = _resolve_src(file_id)
        if not src:
            return self._err("file_id 无效或文件不存在")

        sheets_cfg = data.get("sheets", [])
        colmaps: dict[str, ColumnMap] = {}
        enabled_sheets = []
        for s in sheets_cfg:
            if not s.get("enabled"):
                continue
            name = s["name"]
            cm = s.get("colmap", {})
            colmaps[name] = ColumnMap(
                cert_no=cm.get("cert_no", "I") or "I",
                result=cm.get("result", "M") or "M",
                manufacturer=cm.get("manufacturer", ""),
                factory=cm.get("factory", ""),
                model=cm.get("model", ""),
                standard=cm.get("standard", ""),
                cert_status=cm.get("cert_status", ""),
                agency=cm.get("agency", ""),
                first_data_row=int(cm.get("first_data_row", 2) or 2),
            )
            enabled_sheets.append(name)
        if not enabled_sheets:
            return self._err("未勾选任何工作表")

        try:
            q_interval = float(data.get("query_interval", 1.0))
        except (TypeError, ValueError):
            q_interval = 1.0
        q_interval = max(0.0, min(q_interval, 120.0))     # 限幅 0~120 秒
        cfg_kwargs = {
            "limit": int(data.get("limit", 0) or 0),
            "proxy": data.get("proxy", "") or "",
            "default_agency": data.get("default_agency", "") or "",
            "stop_on_blocked": bool(data.get("stop_on_blocked", True)),
            "write_detail_col": data.get("detail_col", "") or "",
            "row_delay": float(data.get("row_delay", 0.0) or 0.0),
            "query_interval": q_interval,
            "query_jitter": max(0.0, min(float(data.get("query_jitter", 1.0) or 0.0), 60.0)),
        }
        task = _mgr.create_and_start(src, cfg_kwargs, colmaps, enabled_sheets)
        return self._json({"task_id": task.id})

    # ---------------- API: progress ----------------
    def _progress(self, qs):
        task_id = (qs.get("id") or [""])[0]
        offset = int((qs.get("offset") or ["0"])[0])
        task = _mgr.get(task_id)
        if not task:
            return self._err("task 不存在", 404)
        return self._json(task.snapshot(offset))

    # ---------------- API: cancel ----------------
    def _cancel(self):
        data = self._read_json()
        status = _mgr.cancel(data.get("task_id", ""))
        return self._json({"status": status})

    # ---------------- API: download ----------------
    def _download(self, qs):
        task_id = (qs.get("id") or [""])[0]
        kind = (qs.get("kind") or ["diff"])[0]
        task = _mgr.get(task_id)
        if not task:
            return self._err("task 不存在", 404)
        target = task.out_crawl if kind == "crawl" else task.out_xlsx
        out = os.path.normpath(target or "")
        # 只允许 _work 工作目录内的输出文件（防穿越）
        if not out.startswith(os.path.normpath(WORK_DIR)) or not os.path.isfile(out):
            return self._err("结果文件不存在", 404)
        with open(out, "rb") as f:
            body = f.read()
        fn = urllib.parse.quote(os.path.basename(out))
        self.send_response(200)
        self.send_header("Content-Type",
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{fn}")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    # ---------------- API: ul-status / ul-login ----------------
    def _ul_status(self):
        return self._json({
            "logged_in": os.path.isdir(_UL_PROFILE),
            "login_running": _ul_login_running.is_set(),
        })

    def _ul_login(self):
        if _ul_login_running.is_set():
            return self._json({"started": False, "message": "登录已在进行中"})
        _ul_login_running.set()

        def _do():
            try:
                from ..ul_login import login
                login()
            except Exception:
                pass
            finally:
                _ul_login_running.clear()

        threading.Thread(target=_do, daemon=True).start()
        return self._json({"started": True})

    # ---------------- 读取请求体 ----------------
    def _read_json(self) -> dict:
        clen = int(self.headers.get("Content-Length", 0))
        if clen <= 0:
            return {}
        raw = self.rfile.read(clen)
        try:
            return json.loads(raw.decode("utf-8"))
        except Exception:
            return {}


# ---------------- 辅助 ----------------

def _parse_multipart_file(raw: bytes, boundary: bytes):
    """从 multipart body 里取出第一个文件部分，返回 (filename, data)。"""
    delim = b"--" + boundary
    parts = raw.split(delim)
    for part in parts:
        if b"Content-Disposition" not in part:
            continue
        # 头与体以 \r\n\r\n 分隔
        idx = part.find(b"\r\n\r\n")
        if idx < 0:
            continue
        header = part[:idx].decode("utf-8", "replace")
        if "filename=" not in header:
            continue
        m = re.search(r'filename="([^"]*)"', header)
        filename = m.group(1) if m else "upload.xlsx"
        body = part[idx + 4:]
        # 去掉结尾的 \r\n
        if body.endswith(b"\r\n"):
            body = body[:-2]
        return filename, body
    return "", b""


def _read_sheets_preview(path: str, header_rows: int = 3) -> list:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            preview = []
            header_vals: dict[int, str] = {}
            for r, row in enumerate(ws.iter_rows(min_row=1, max_row=header_rows,
                                                 values_only=True), 1):
                cells = []
                for c, v in enumerate(row, 1):
                    if r == 1 and v is not None and str(v).strip() != "":
                        header_vals[c] = str(v).strip()[:40]
                    if v is None or str(v).strip() == "":
                        continue
                    cells.append({"col": get_column_letter(c),
                                  "value": str(v)[:60]})
                preview.append({"row": r, "cells": cells})
            # 全列清单（供前端下拉选列用）：以第 1 行为表头，含空表头列
            columns = []
            for c in range(1, (ws.max_column or 0) + 1):
                columns.append({"col": get_column_letter(c),
                                "name": header_vals.get(c, "")})
            out.append({"name": ws.title, "max_row": ws.max_row,
                        "max_col": ws.max_column, "preview": preview,
                        "columns": columns})
    finally:
        wb.close()
    return out


def _resolve_src(file_id: str) -> str:
    """由 file_id 定位上传的 xlsx，路径校验防穿越。"""
    if not file_id or not re.fullmatch(r"[0-9a-f]{6,}", file_id):
        return ""
    d = os.path.normpath(os.path.join(WORK_DIR, file_id))
    if not d.startswith(os.path.normpath(WORK_DIR)) or not os.path.isdir(d):
        return ""
    for fn in os.listdir(d):
        if fn.lower().endswith(".xlsx") and not fn.endswith("_核对结果.xlsx"):
            return os.path.join(d, fn)
    return ""


def serve(host: str = "127.0.0.1", port: int = 8756, open_browser: bool = True) -> None:
    os.makedirs(WORK_DIR, exist_ok=True)
    httpd = None
    for p in range(port, port + 20):
        try:
            httpd = ThreadingHTTPServer((host, p), Handler)
            port = p
            break
        except OSError:
            continue
    if httpd is None:
        raise SystemExit(f"端口 {port}~{port + 19} 均被占用，无法启动")

    url = f"http://{host}:{port}/"
    print("=" * 56)
    print(f"  证书核对网页已启动： {url}")
    print("  如果浏览器没自动打开，请手动复制上面地址到浏览器。")
    print("  ★ 使用期间请勿关闭本窗口（关闭窗口 = 停止服务）。")
    print("  停止服务：在本窗口按 Ctrl+C，或直接关闭本窗口。")
    print("=" * 56)

    if open_browser:
        # 放到后台线程里、等服务开始接受连接后再打开，避免“打开太早 → 拒绝连接”
        def _open():
            import time
            time.sleep(1.0)
            try:
                import webbrowser
                if webbrowser.open(url):
                    return
            except Exception:
                pass
            try:                       # 兜底：用系统默认方式打开
                os.startfile(url)      # type: ignore[attr-defined]
            except Exception:
                pass
        threading.Thread(target=_open, daemon=True).start()

    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n已停止。")
        httpd.shutdown()
