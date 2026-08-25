"""核对内核 —— Web 与 CLI 共用。

把「逐行核对」从 excel_runner 的一次性同步流程里抽出来，使其：
  - 惰性逐行产出结构化结果（RowOutcome），调用方可在每行后上报进度、检查取消；
  - 字段级结论（FieldVerdict）由 compare.compare_fields 产出，是 Web 展示、
    M 列文案、Excel 标红三者共同的唯一判断来源；
  - 携带每字段所在列字母，供 excel_writer 精确标红。

不发起任何写文件动作；读 xlsx 取值 + 查询 + 比对，产出结果交给上层。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Iterator

from .compare import FieldVerdict, compare_fields, summarize_fields
from .config import ColumnMap, RunConfig
from .models import BLOCKED, FOUND, NOT_FOUND, CertRecord
from .providers import ProviderRouter


def col_letter_to_idx(letter: str) -> int:
    """'A'->1, 'I'->9, 'M'->13；空字符串->0。"""
    n = 0
    for ch in (letter or "").strip().upper():
        n = n * 26 + (ord(ch) - 64)
    return n


@dataclass
class FieldCell:
    """字段级结论 + 它在表格中的列字母（用于标红定位）。"""

    verdict: FieldVerdict
    column: str           # 该字段映射的列字母；空表示表格无此列，不可标红


@dataclass
class RowOutcome:
    """一行核对的完整结果。"""

    sheet: str
    row: int
    cert_no: str
    agency: str                       # 实际路由到的 Provider 名（rec.source）
    status: str                       # CertRecord.status
    consistent: bool                  # 无 MISMATCH 即 True（非 found 时为 False）
    verdict_text: str                 # 写入 M 列的摘要（含官网值）
    fields: list[FieldCell] = field(default_factory=list)
    source_url: str = ""
    message: str = ""
    result_col: str = "M"             # 结论写入列
    detail_col: str = ""              # 详情（URL/错误）写入列

    def to_json(self) -> dict:
        return {
            "sheet": self.sheet,
            "row": self.row,
            "cert_no": self.cert_no,
            "agency": self.agency,
            "status": self.status,
            "consistent": self.consistent,
            "verdict_text": self.verdict_text,
            "source_url": self.source_url,
            "message": self.message,
            "fields": [
                {
                    "field": fc.verdict.field,
                    "label": fc.verdict.label,
                    "sheet_value": fc.verdict.sheet_value,
                    "site_value": fc.verdict.site_value,
                    "verdict": fc.verdict.verdict,
                    "detail": fc.verdict.detail,
                    "column": fc.column,
                }
                for fc in self.fields
            ],
        }


class VerifyEngine:
    """惰性逐行核对引擎。

    用法：
        eng = VerifyEngine(cfg)
        for outcome in eng.iter_rows():
            ...  # 上报进度 / 检查取消 / 收集 outcome
        # 之后由 excel_writer 落盘

    total 属性在首次 iter_rows 前需调用 count_rows() 预估（Web 用于进度条）。
    """

    def __init__(self, cfg: RunConfig, router: ProviderRouter | None = None):
        self.cfg = cfg
        self.router = router or ProviderRouter(
            headless=cfg.headless, proxy=cfg.proxy or None,
            interval=getattr(cfg, "query_interval", None),
            jitter=getattr(cfg, "query_jitter", None))
        self._total = 0

    # ---------- 预估总条数（供进度条） ----------
    def count_rows(self) -> int:
        from openpyxl import load_workbook

        wb = load_workbook(self.cfg.xlsx_path, read_only=True, data_only=True)
        total = 0
        try:
            sheets = self.cfg.sheets or wb.sheetnames
            for name in sheets:
                if name not in wb.sheetnames:
                    continue
                ws = wb[name]
                cmap = self.cfg.column_map_for(name)
                col = col_letter_to_idx(cmap.cert_no)
                cnt = 0
                for r in range(cmap.first_data_row, ws.max_row + 1):
                    v = ws.cell(row=r, column=col).value
                    if v is not None and str(v).strip():
                        cnt += 1
                        if self.cfg.limit and total + cnt >= self.cfg.limit:
                            break
                total += cnt
                if self.cfg.limit and total >= self.cfg.limit:
                    total = self.cfg.limit
                    break
        finally:
            wb.close()
        self._total = total
        return total

    @property
    def total(self) -> int:
        return self._total

    # ---------- 逐行核对 ----------
    def iter_rows(self) -> Iterator[RowOutcome]:
        from openpyxl import load_workbook

        wb = load_workbook(self.cfg.xlsx_path, read_only=True, data_only=True)
        produced = 0
        try:
            sheets = self.cfg.sheets or wb.sheetnames
            for name in sheets:
                if name not in wb.sheetnames:
                    continue
                ws = wb[name]
                cmap: ColumnMap = self.cfg.column_map_for(name)

                for row in range(cmap.first_data_row, ws.max_row + 1):
                    if self.cfg.limit and produced >= self.cfg.limit:
                        return
                    cert_no = self._cell(ws, row, cmap.cert_no)
                    if not cert_no:
                        continue

                    outcome = self._verify_row(ws, name, row, cert_no, cmap)
                    produced += 1
                    yield outcome

                    if outcome.status == BLOCKED and self.cfg.stop_on_blocked:
                        return
        finally:
            wb.close()
            self._close_router()

    def _close_router(self) -> None:
        close = getattr(self.router, "close", None)
        if callable(close):
            try:
                close()
            except Exception:
                pass

    # ---------- 单行 ----------
    def _verify_row(self, ws, sheet: str, row: int, cert_no: str,
                    cmap: ColumnMap) -> RowOutcome:
        sheet_row = {
            "manufacturer": self._cell(ws, row, cmap.manufacturer),
            "factory": self._cell(ws, row, cmap.factory),
            "model": self._cell(ws, row, cmap.model),
            "standard": self._cell(ws, row, cmap.standard),
            "cert_status": self._cell(ws, row, cmap.cert_status),
        }
        agency = self._cell(ws, row, cmap.agency) or self.cfg.default_agency
        rec = self.router.query(cert_no, agency)
        return self.build_outcome(sheet, row, cert_no, rec, sheet_row, cmap,
                                  self.cfg.write_detail_col)

    @staticmethod
    def build_outcome(sheet: str, row: int, cert_no: str, rec: CertRecord,
                      sheet_row: dict, cmap: ColumnMap,
                      detail_col: str = "") -> RowOutcome:
        """由 CertRecord + 表格值组装 RowOutcome（可离线单测）。"""
        from .compare import compare_record

        fvs = compare_fields(sheet_row, rec)
        col_of = {
            "manufacturer": cmap.manufacturer,
            "factory": cmap.factory,
            "model": cmap.model,
            "standard": cmap.standard,
            "cert_status": cmap.cert_status,
        }
        field_cells = [FieldCell(verdict=fv, column=col_of.get(fv.field, "")) for fv in fvs]

        if rec.status == FOUND:
            consistent, verdict_text = summarize_fields(fvs)
        else:
            consistent, verdict_text = compare_record(sheet_row, rec)

        return RowOutcome(
            sheet=sheet, row=row, cert_no=cert_no,
            agency=rec.source, status=rec.status,
            consistent=consistent, verdict_text=verdict_text,
            fields=field_cells,
            source_url=rec.source_url, message=rec.message,
            result_col=cmap.result, detail_col=detail_col,
        )

    @staticmethod
    def _cell(ws, row: int, col_letter: str) -> str:
        if not col_letter:
            return ""
        v = ws.cell(row=row, column=col_letter_to_idx(col_letter)).value
        return "" if v is None else str(v).strip()


def summarize_outcomes(outcomes: list[RowOutcome]) -> dict:
    """累计统计，供进度/汇总用。"""
    total = len(outcomes)
    found = sum(1 for o in outcomes if o.status == FOUND)
    nf = sum(1 for o in outcomes if o.status == NOT_FOUND)
    blk = sum(1 for o in outcomes if o.status == BLOCKED)
    consistent = sum(1 for o in outcomes if o.consistent)
    return {
        "total": total,
        "found": found,
        "not_found": nf,
        "blocked": blk,
        "consistent": consistent,
        "review": total - consistent,
    }
