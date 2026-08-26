"""读 xlsx（多 sheet）→ 按 I 列证书号查官网 → 比对 → 写 M 列。

列映射用 config.py 里的 ColumnMap 配置。默认只锁定了题目明确给的两列：
    I 列 = 证书号（输入）
    M 列 = 比对结果（输出）
其余五个待比对字段所在列，先用 `python -m cert_verify.cli inspect <xlsx>`
打印表头，再把列号填进 config.py。
"""
from __future__ import annotations

import os
import time
from dataclasses import dataclass

from .compare import compare_record
from .config import ColumnMap, RunConfig
from .models import BLOCKED, CertRecord
from .providers import ProviderRouter


@dataclass
class RowResult:
    sheet: str
    row: int
    cert_no: str
    verdict: str
    consistent: bool
    record: CertRecord


def col_letter_to_idx(letter: str) -> int:
    """'A'->1, 'I'->9, 'M'->13"""
    n = 0
    for ch in letter.strip().upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def inspect_workbook(path: str, header_rows: int = 3) -> str:
    """打印每个 sheet 的前几行，用于确认列映射。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = [f"文件: {path}", f"Sheet 列表: {wb.sheetnames}", ""]
    for ws in wb.worksheets:
        out.append(f"=== [{ws.title}]  {ws.max_row} 行 × {ws.max_column} 列 ===")
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=header_rows, values_only=True), 1):
            for c, v in enumerate(row, 1):
                if v is None or str(v).strip() == "":
                    continue
                from openpyxl.utils import get_column_letter
                out.append(f"  r{r} {get_column_letter(c):>2}({c:>2}) = {str(v)[:60]!r}")
            out.append("  " + "-" * 40)
    wb.close()
    return "\n".join(out)


def _cell(ws, row: int, col_letter: str) -> str:
    if not col_letter:
        return ""
    v = ws.cell(row=row, column=col_letter_to_idx(col_letter)).value
    return "" if v is None else str(v).strip()


def run(cfg: RunConfig) -> list[RowResult]:
    """读 xlsx → 逐行核对 → 比对 → 写 M 列（并对真不一致字段标红）。

    改为基于 VerifyEngine（Web/CLI 共用内核）+ excel_writer（标红写入），
    行为与旧版一致：输出到副本、stop_on_blocked、逐行打印、summarize。
    """
    from .excel_writer import write_results
    from .verifier import RowOutcome, VerifyEngine, summarize_outcomes  # noqa: F401

    engine = VerifyEngine(cfg)
    results: list[RowResult] = []
    outcomes_by_sheet: dict[str, list] = {}
    cur_sheet = None

    for o in engine.iter_rows():
        if o.sheet != cur_sheet:
            cur_sheet = o.sheet
            print(f"\n=== Sheet [{o.sheet}] ===")
        outcomes_by_sheet.setdefault(o.sheet, []).append(o)
        results.append(RowResult(o.sheet, o.row, o.cert_no, o.verdict_text,
                                 o.consistent, _rec_stub(o)))
        print(f"  r{o.row:<4} {o.cert_no:<18} {o.agency:<14} {o.verdict_text[:90]}")

        if o.status == BLOCKED and cfg.stop_on_blocked:
            print("  [stop] 检测到网络拦截，先解决网络再继续，避免整表写满错误结果。")
            break

        time.sleep(cfg.row_delay)

    out = cfg.output_path
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    info = write_results(cfg.xlsx_path, out, outcomes_by_sheet,
                         detail_col=cfg.write_detail_col)
    if info.get("degraded"):
        print(f"[warn] {info.get('message')}")
    print(f"\n已写入: {out}")
    return results


def _rec_stub(outcome) -> CertRecord:
    """从 RowOutcome 还原一个精简 CertRecord，供 summarize 统计状态用。"""
    rec = CertRecord(query_no=outcome.cert_no, source=outcome.agency,
                     status=outcome.status, message=outcome.message,
                     source_url=outcome.source_url)
    return rec


def summarize(results: list[RowResult]) -> str:
    total = len(results)
    ok = sum(1 for r in results if r.consistent)
    from .models import BLOCKED, FOUND, NOT_FOUND
    found = sum(1 for r in results if r.record.status == FOUND)
    nf = sum(1 for r in results if r.record.status == NOT_FOUND)
    blk = sum(1 for r in results if r.record.status == BLOCKED)
    return (f"共 {total} 条 | 查到 {found} | 官网无 {nf} | 被拦截 {blk} | "
            f"字段全一致 {ok} | 需人工复核 {total - ok}")
