"""Excel 结果写入层：在原文件副本上写结论 + 对真不一致字段标红字。

约束（对应设计正确性属性）：
  - 绝不写回源文件：load 源 → 另存 out。
  - 只有字段 verdict == MISMATCH 且该字段有列映射时，才把对应单元格设红色字体。
  - 红色字体保留原字体其他属性（字号、加粗、字体名），不动填充与边框。
  - MISSING_SITE / MISSING_SHEET / 非 found 状态不改字段样式。
  - 保存异常时降级为「仅写值不设样式」再存一次。
"""
from __future__ import annotations

from .compare import MISMATCH, MISSING_SITE
from .models import FOUND
from .verifier import RowOutcome, col_letter_to_idx

RED = "FFFF0000"
BLUE = "FF00B0F0"      # 天蓝色：抓取表里“官网抓取值”用天蓝字，和表格原值区分


def _red_font(cell):
    """复制单元格原字体，仅把颜色改成红色。"""
    from openpyxl.styles import Font

    old = cell.font
    return Font(
        name=old.name, size=old.size, bold=old.bold, italic=old.italic,
        vertAlign=old.vertAlign, underline=old.underline, strike=old.strike,
        color=RED,
    )


def _apply_row(ws, o: RowOutcome, detail_col: str, mark_red: bool) -> None:
    # 结论写 M 列
    if o.result_col:
        ws.cell(row=o.row, column=col_letter_to_idx(o.result_col)).value = o.verdict_text
    # 详情列（URL / 错误）
    col = detail_col or o.detail_col
    if col:
        ws.cell(row=o.row, column=col_letter_to_idx(col)).value = o.source_url or o.message
    # 标红：仅 MISMATCH 且有列映射
    if mark_red:
        for fc in o.fields:
            if fc.verdict.verdict == MISMATCH and fc.column:
                cell = ws.cell(row=o.row, column=col_letter_to_idx(fc.column))
                cell.font = _red_font(cell)


def _colored_font(cell, rgb):
    from openpyxl.styles import Font
    old = cell.font
    return Font(name=old.name, size=old.size, bold=old.bold, italic=old.italic,
                vertAlign=old.vertAlign, underline=old.underline, strike=old.strike,
                color=rgb)


def write_crawl_copy(src_xlsx: str, out_xlsx: str,
                     outcomes_by_sheet: dict[str, list[RowOutcome]],
                     detail_col: str = "N") -> dict:
    """抓取结果导出（沿用原表布局）：把 制造商/型号/标准/证书状态 列的值
    **替换为官网抓取到的值**（蓝字，与表格原值区分），M 列写查询状态、
    详情列写来源 URL。网站无此项写“(官网未提供)”。原文件不动。
    """
    from openpyxl import load_workbook

    STATUS_ZH = {"found": "查到", "not_found": "官网无", "blocked": "被拦截",
                 "error": "查询失败", "unsupported": "未支持"}

    def _write(colorize: bool):
        wb = load_workbook(src_xlsx)
        try:
            for sheet, outcomes in outcomes_by_sheet.items():
                if sheet not in wb.sheetnames:
                    continue
                ws = wb[sheet]
                for o in outcomes:
                    # 各可比字段：把官网值写回对应列
                    for fc in o.fields:
                        if not fc.column:
                            continue
                        if fc.verdict.verdict == MISSING_SITE:
                            val = "(官网未提供)"
                        else:
                            val = fc.verdict.site_value
                        cell = ws.cell(row=o.row, column=col_letter_to_idx(fc.column))
                        cell.value = val
                        if colorize and val and val != "(官网未提供)":
                            cell.font = _colored_font(cell, BLUE)
                    # M 列写查询状态；详情列写来源 URL
                    if o.result_col:
                        ws.cell(row=o.row, column=col_letter_to_idx(o.result_col)).value = (
                            STATUS_ZH.get(o.status, o.status)
                            + (("：" + o.message) if o.status != FOUND and o.message else ""))
                    col = detail_col or o.detail_col
                    if col:
                        ws.cell(row=o.row, column=col_letter_to_idx(col)).value = (
                            o.source_url or o.message)
            wb.save(out_xlsx)
        finally:
            wb.close()

    try:
        _write(colorize=True)
        return {"degraded": False, "message": ""}
    except Exception as e:  # noqa: BLE001
        try:
            _write(colorize=False)
            return {"degraded": True, "message": f"着色失败已降级：{type(e).__name__}: {e}"}
        except Exception as e2:  # noqa: BLE001
            return {"degraded": True, "message": f"写入失败：{type(e2).__name__}: {e2}"}


def write_results(src_xlsx: str, out_xlsx: str,
                  outcomes_by_sheet: dict[str, list[RowOutcome]],
                  detail_col: str = "") -> dict:
    """把逐行结果写进源文件副本并另存 out_xlsx。

    返回 {"degraded": bool, "message": str}。
    """
    from openpyxl import load_workbook

    def _write(mark_red: bool):
        wb = load_workbook(src_xlsx)          # 保留样式/公式/合并；不 read_only
        try:
            for sheet, outcomes in outcomes_by_sheet.items():
                if sheet not in wb.sheetnames:
                    continue
                ws = wb[sheet]
                for o in outcomes:
                    _apply_row(ws, o, detail_col, mark_red)
            wb.save(out_xlsx)
        finally:
            wb.close()

    try:
        _write(mark_red=True)
        return {"degraded": False, "message": ""}
    except Exception as e:  # noqa: BLE001
        # 降级：不设样式，只写值，尽量还是给出可用结果
        try:
            _write(mark_red=False)
            return {"degraded": True,
                    "message": f"标红失败已降级为仅写值：{type(e).__name__}: {e}"}
        except Exception as e2:  # noqa: BLE001
            return {"degraded": True,
                    "message": f"写入失败：{type(e2).__name__}: {e2}"}
