"""excel_writer 单元测试（离线）：

  - 仅 MISMATCH 字段单元格字体变红，原字体属性保留；
  - MISSING/一致字段不变红；
  - 源文件字节哈希写入前后不变（原文件不可变）。
"""
import hashlib
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from cert_verify.compare import compare_fields  # noqa: E402
from cert_verify.config import ColumnMap  # noqa: E402
from cert_verify.excel_writer import RED, write_results  # noqa: E402
from cert_verify.models import FOUND, CertRecord  # noqa: E402
from cert_verify.verifier import VerifyEngine  # noqa: E402

_fail = 0


def check(name, cond):
    global _fail
    print(f"  [{'OK ' if cond else 'FAIL'}] {name}")
    if not cond:
        _fail += 1


def _sha1(path):
    with open(path, "rb") as f:
        return hashlib.sha1(f.read()).hexdigest()


def build_src(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    # E=型号 G=制造商 H=标准 I=证书号 M=结果
    ws["E1"] = "型号"; ws["G1"] = "制造商"; ws["H1"] = "标准"; ws["I1"] = "证书号"
    ws["E2"] = "A24-999XXX"        # 与官网 A24-065N3A 不一致 → 应标红
    ws["G2"] = "Delta Electronics" # 与官网一致 → 不标红
    ws["H2"] = "EN 62368-1"
    ws["I2"] = "R50151509"
    # 给 E2 一个非默认字体，验证标红保留其属性
    ws["E2"].font = Font(name="Consolas", size=14, bold=True)
    wb.save(path)


cmap = ColumnMap(cert_no="I", result="M", manufacturer="G", model="E",
                 standard="H", cert_status="", factory="", first_data_row=2)

tmp = tempfile.mkdtemp()
src = os.path.join(tmp, "src.xlsx")
out = os.path.join(tmp, "out.xlsx")
build_src(src)
src_hash_before = _sha1(src)

# 构造一条 FOUND 记录：型号不一致、制造商一致、标准一致
rec = CertRecord(query_no="R50151509", source="TUV-Rheinland", status=FOUND,
                 manufacturer="Delta Electronics Inc.", model="A24-065N3A",
                 standard="EN 62368-1", cert_status="Valid Certificate",
                 source_url="https://certipedia.com/x")
row = {"manufacturer": "Delta Electronics", "model": "A24-999XXX",
       "standard": "EN 62368-1", "factory": "", "cert_status": ""}
outcome = VerifyEngine.build_outcome("S1", 2, "R50151509", rec, row, cmap, detail_col="N")

info = write_results(src, out, {"S1": [outcome]}, detail_col="N")
check("写入未降级", info["degraded"] is False)

from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(out)
ws = wb["S1"]
e2 = ws["E2"]; g2 = ws["G2"]
check("型号单元格字体变红", (e2.font.color is not None and e2.font.color.rgb == RED))
check("制造商单元格未变红", (g2.font.color is None or g2.font.color.rgb != RED))
check("标红保留原字体名", e2.font.name == "Consolas")
check("标红保留原字号", e2.font.size == 14)
check("标红保留加粗", e2.font.bold is True)
check("M 列写入结论", ws["M2"].value and "不一致" in ws["M2"].value)
check("N 列写入来源URL", ws["N2"].value == "https://certipedia.com/x")
wb.close()

check("源文件字节未改", _sha1(src) == src_hash_before)

print("\n" + ("全部通过" if _fail == 0 else f"{_fail} 项失败"))
sys.exit(1 if _fail else 0)
